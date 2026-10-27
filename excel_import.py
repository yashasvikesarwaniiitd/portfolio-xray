"""Read an Excel workbook as portfolio text, so everything downstream stays CSV.

The whole product speaks CSV text: merging, saving, analysis and the chat all take it. So
Excel support is a translation at the door rather than a second path through the system -
a workbook is turned into the same text a CSV export would have produced, and from there it
is indistinguishable from one. That is why this adds no new dependency (openpyxl already
writes the Excel report) and no new code in the merge, storage or analysis layers.

SHEETS. A broker workbook is rarely just the data: there is a cover sheet, a summary, a
second tab for another asset class. Each sheet is tried on its own merits - it counts if
portfolio_merge can read it as a transaction statement or a SIP sheet - and every one that
qualifies becomes a part of the portfolio. Sheets that don't qualify are listed with a
reason rather than silently dropped, because "your workbook had 3 sheets and I read 1" is
information the person needs.

THE EXCEL-SPECIFIC TRAP. Excel "helpfully" turns the header `Mar'25` into the date 1-Mar-2025.
Read back as a date object it is no longer a month label, so the SIP parser would see a
sheet with no month columns and reconstruct nothing - silently, from a file that looks
identical on screen. On a SIP header row, date cells are therefore written back as `Mar'25`.
Everywhere else a date becomes an ISO date, which the statement reader already understands.

BOUNDS. A workbook is a zip, and a hostile or simply enormous one should cost a rejected
upload rather than the server's memory: sheets and rows are capped, and the file is read in
openpyxl's streaming mode. Only .xlsx and .xlsm are supported; the legacy .xls format is a
different binary and is refused with a message that says how to convert it.
"""
from __future__ import annotations

import csv
import io
from datetime import date, datetime

from portfolio_merge import PART_DELIM, SIP, parse_part

MAX_SHEETS = 30
MAX_ROWS_PER_SHEET = 20_000
MAX_BYTES = 5_000_000


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e15:
            return str(int(v))
        return format(v, ".10f").rstrip("0").rstrip(".")
    return " ".join(str(v).split())


def _is_sip_header(row) -> bool:
    names = {str(c).strip().lower() for c in row if isinstance(c, str)}
    return "where" in names and "total invested" in names


def _sheet_to_text(rows: list) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    for row in rows:
        if _is_sip_header(row):
            cells = [v.strftime("%b'%y") if isinstance(v, (datetime, date)) else _cell(v)
                     for v in row]
        else:
            cells = [_cell(v) for v in row]
        while cells and cells[-1] == "":
            cells.pop()
        if any(cells):
            w.writerow(cells)
    return buf.getvalue()


def is_excel(name: str, data: bytes = b"") -> bool:
    """By extension, or by the zip signature when the name is unhelpful."""
    n = (name or "").lower()
    return n.endswith((".xlsx", ".xlsm")) or (data[:2] == b"PK" and not n.endswith(".csv"))


def workbook_to_text(data: bytes) -> tuple:
    """(portfolio_text, report). Raises ValueError with a reason the person can act on."""
    if len(data) > MAX_BYTES:
        raise ValueError(f"That workbook is {len(data) // 1_000_000} MB; the limit is "
                         f"{MAX_BYTES // 1_000_000} MB.")
    if data[:2] != b"PK":
        raise ValueError("That is not an .xlsx workbook. If it is an old .xls file, open it "
                         "in Excel and use Save As > Excel Workbook (.xlsx) or CSV.")
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"Could not open that file as an Excel workbook ({e}).")

    used, ignored, parts = [], [], []
    try:
        for ws in wb.worksheets[:MAX_SHEETS]:
            if getattr(ws, "sheet_state", "visible") != "visible":
                ignored.append({"sheet": ws.title, "reason": "hidden sheet"})
                continue
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_ROWS_PER_SHEET:
                    break
                rows.append(row)
            text = _sheet_to_text(rows)
            if not text.strip():
                ignored.append({"sheet": ws.title, "reason": "empty"})
                continue
            try:
                part = parse_part(text)
            except ValueError:
                first = next((r for r in text.splitlines() if r.strip()), "")
                ignored.append({"sheet": ws.title,
                                "reason": "not a transaction statement or SIP sheet "
                                          f"(starts: {first[:60]!r})"})
                continue
            if not part.rows:
                ignored.append({"sheet": ws.title, "reason": "a header but no rows"})
                continue
            parts.append(text)
            used.append({"sheet": ws.title, "kind": part.kind, "rows": len(part.rows)})
        if len(wb.worksheets) > MAX_SHEETS:
            ignored.append({"sheet": f"{len(wb.worksheets) - MAX_SHEETS} more",
                            "reason": f"only the first {MAX_SHEETS} sheets are read"})
    finally:
        wb.close()

    if not parts:
        seen = ", ".join(f"{i['sheet']} ({i['reason']})" for i in ignored) or "no sheets"
        raise ValueError("No sheet in that workbook looks like a transaction statement or a "
                         "SIP sheet. A statement needs a date column and a quantity column; "
                         f"a SIP sheet needs 'Where' and 'Total Invested'. Sheets seen: {seen}.")
    text = parts[0] if len(parts) == 1 else ("\n" + PART_DELIM + "\n").join(parts) + "\n"
    return text, {"sheets_used": used, "sheets_ignored": ignored}


def describe(report: dict) -> str:
    used = report["sheets_used"]
    msg = "Read " + ", ".join(f"“{u['sheet']}” ({u['rows']} rows)" for u in used) + "."
    if report["sheets_ignored"]:
        msg += " Skipped: " + ", ".join(f"“{i['sheet']}”" for i in report["sheets_ignored"]) + "."
    return msg
