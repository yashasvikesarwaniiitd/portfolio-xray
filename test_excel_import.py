"""Tests for reading Excel workbooks as portfolios.

Workbooks are built in memory with openpyxl, so these are hermetic. The point of the module
is that an .xlsx becomes the SAME text a CSV export would have been, so several tests prove
that by feeding the result through the real readers and the real merge rather than
inspecting the text.
"""
import io
from datetime import datetime

import pytest
from openpyxl import Workbook

import agent
import excel_import as X
import portfolio_merge as M


def book(**sheets):
    """{sheet name: rows} -> .xlsx bytes."""
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


STMT_ROWS = [
    ["Groww Order Statement"],
    ["Account", "demo"],
    [],
    ["Date", "Symbol", "Stock name", "Type", "Quantity", "Price", "Order status"],
    [datetime(2026, 1, 5, 9, 15), "INFY", "Infosys", "BUY", 10, 1500.0, "Executed"],
    [datetime(2026, 2, 10, 14, 2), "HDFCBANK", "HDFC Bank", "BUY", 5, 1600.5, "Executed"],
]
STMT_CSV = ("Date,Symbol,Stock name,Type,Quantity,Price,Order status\n"
            "2026-01-05,INFY,Infosys,BUY,10,1500,Executed\n"
            "2026-02-10,HDFCBANK,HDFC Bank,BUY,5,1600.5,Executed\n")


def holdings(text, tmp_path):
    p = tmp_path / "p.csv"
    p.write_text(text, encoding="utf-8")
    return agent.read_portfolio(str(p), resolve_funds=False)[0]


def test_a_statement_workbook_reads_like_the_csv_would_have(tmp_path):
    """Real date objects, floats like 1500.0 and banner rows above the header: none of it
    may change what the statement means."""
    text, rep = X.workbook_to_text(book(Orders=STMT_ROWS))
    by = {h["name"]: h for h in holdings(text, tmp_path)}
    assert set(by) == {"Infosys", "HDFC Bank"}
    infy = by["Infosys"]["transactions"][0]
    assert (infy.units, infy.price, infy.date.isoformat()) == (10, 1500.0, "2026-01-05")
    assert rep["sheets_used"] == [{"sheet": "Orders", "kind": "statement", "rows": 2}]


def test_an_excel_upload_after_the_csv_of_the_same_trades_adds_nothing():
    """The property that makes this safe: both routes produce the same row identity, so a
    person who uploaded the CSV and later the .xlsx of one statement is not doubled."""
    text, _ = X.workbook_to_text(book(Orders=STMT_ROWS))
    _, rep = M.merge(STMT_CSV, text)
    assert rep["transactions_added"] == 0 and rep["transactions_duplicate"] == 2


def test_a_sip_header_that_excel_turned_into_dates_is_still_a_month_column(tmp_path):
    """Excel converts Mar'25 into 1-Mar-2025 on entry. Read back as a date it is no longer
    a month label, and the SIP parser would find no months and reconstruct nothing - from a
    file that looks identical on screen."""
    rows = [["Mode", "App", "Type", "Market", "Risk", "Where", "symbol", "source",
             "Estimated Returns (3Y)", "Total Invested",
             datetime(2025, 3, 1), datetime(2025, 4, 1)],
            ["Equity", "Groww", "Share", "India", "Med", "Reliance", "RELIANCE.NS",
             "yfinance", "-", 10000, 5000, 5000]]
    text, rep = X.workbook_to_text(book(SIP=rows))
    assert "Mar'25" in text and "Apr'25" in text
    h = holdings(text, tmp_path)[0]
    assert sorted((y, m) for y, m, _ in h["inflows"]) == [(2025, 3), (2025, 4)]
    assert rep["sheets_used"][0]["kind"] == "sip"


def test_dates_outside_a_sip_header_become_iso_dates_not_month_labels():
    text, _ = X.workbook_to_text(book(Orders=STMT_ROWS))
    assert "2026-01-05" in text and "'26" not in text


def test_a_cover_sheet_is_skipped_with_a_reason_and_the_data_sheet_is_read():
    text, rep = X.workbook_to_text(book(
        Cover=[["Portfolio report"], ["Generated 2026-10-01"]], Orders=STMT_ROWS))
    assert [u["sheet"] for u in rep["sheets_used"]] == ["Orders"]
    assert rep["sheets_ignored"][0]["sheet"] == "Cover"
    assert "not a transaction statement" in rep["sheets_ignored"][0]["reason"]


def test_two_data_sheets_become_two_sections_of_one_portfolio(tmp_path):
    """A workbook with stocks on one tab and funds on another: the multi-part portfolio
    model exists for exactly this."""
    mf = [["Scheme Name", "Transaction Type", "Units", "NAV", "Date", "Amount"],
          ["Parag Parikh Flexi Cap Fund - Direct Growth", "Purchase", 10, 70,
           datetime(2026, 1, 15), 700]]
    text, rep = X.workbook_to_text(book(Stocks=STMT_ROWS, Funds=mf))
    assert len(M.split_parts(text)) == 2 and len(rep["sheets_used"]) == 2
    names = {h["name"] for h in holdings(text, tmp_path)}
    assert {"Infosys", "HDFC Bank", "Parag Parikh Flexi Cap Fund - Direct Growth"} <= names


def test_a_hidden_sheet_is_not_read():
    wb = Workbook()
    wb.active.title = "Orders"
    for r in STMT_ROWS:
        wb.active.append(r)
    hidden = wb.create_sheet("Old")
    for r in STMT_ROWS:
        hidden.append(r)
    hidden.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    _, rep = X.workbook_to_text(buf.getvalue())
    assert [u["sheet"] for u in rep["sheets_used"]] == ["Orders"]
    assert rep["sheets_ignored"] == [{"sheet": "Old", "reason": "hidden sheet"}]


def test_a_workbook_with_nothing_usable_says_what_it_saw():
    with pytest.raises(ValueError, match="Sheets seen: .*Notes"):
        X.workbook_to_text(book(Notes=[["just"], ["some text"]]))


def test_an_old_xls_or_a_renamed_text_file_is_refused_with_how_to_fix_it():
    with pytest.raises(ValueError, match="Save As"):
        X.workbook_to_text(b"\xd0\xcf\x11\xe0 legacy xls bytes")
    with pytest.raises(ValueError, match="not an .xlsx"):
        X.workbook_to_text(b"a,b,c\n1,2,3\n")


def test_a_corrupt_zip_is_a_clean_error_not_a_crash():
    with pytest.raises(ValueError, match="Could not open"):
        X.workbook_to_text(b"PK\x03\x04 this is not really a workbook")


def test_an_oversized_workbook_is_refused_before_it_is_opened(monkeypatch):
    monkeypatch.setattr(X, "MAX_BYTES", 10)
    with pytest.raises(ValueError, match="limit"):
        X.workbook_to_text(b"PK" + b"x" * 50)


def test_numbers_are_written_without_float_noise():
    assert X._cell(1500.0) == "1500"
    assert X._cell(3.5) == "3.5"
    assert X._cell(0.1 + 0.2) == "0.3"
    assert X._cell(None) == "" and X._cell(12) == "12"


def test_extension_detection():
    assert X.is_excel("orders.xlsx") and X.is_excel("A.XLSM")
    assert not X.is_excel("orders.csv")
    assert X.is_excel("download", b"PK\x03\x04...")
    assert not X.is_excel("orders.csv", b"PK\x03\x04...")


def test_describe_names_what_was_read_and_what_was_skipped():
    _, rep = X.workbook_to_text(book(Cover=[["x"]], Orders=STMT_ROWS))
    d = X.describe(rep)
    assert "“Orders” (2 rows)" in d and "Skipped: “Cover”" in d
