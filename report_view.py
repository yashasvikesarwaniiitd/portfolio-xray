"""Read a built health report back out, so the app can show what is IN it.

The report is an Excel workbook, and the app used to treat it as an opaque download: a button
and a caption, with no way to see what a report said without saving it and opening Excel. The
Health report tab shows the executive summary in the page instead. The workbook is the only
copy of those findings - the AI sentences are written into it and nowhere else - so they are
read back from it.

This is deliberately tolerant. It reads cells by their content, not their coordinates, so a
restyle that moves a row does not break it, and any failure (a corrupt file, a sheet that has
been renamed) returns "nothing found" rather than raising: a preview must never stop someone
downloading the report it previews. Pure and dependency-light (openpyxl only), so the
frontend can import it without touching any backend module.
"""
from __future__ import annotations

import io
import re

SUMMARY_SHEET = "Executive Summary"
# Sheets that explain the report rather than analyse the portfolio.
_META = ("How To Use", "Methodology")


def _open(data: bytes):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(data), read_only=True, data_only=True)


def findings(data: bytes, limit: int = 8) -> list:
    """The executive-summary bullets, without their bullet character."""
    try:
        wb = _open(data)
        try:
            ws = wb[SUMMARY_SHEET]
            out = []
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if isinstance(v, str) and v.strip().startswith("•"):
                        out.append(v.strip().lstrip("•").strip())
            return out[:limit]
        finally:
            wb.close()
    except Exception:
        return []


def score(data: bytes):
    """The diversification score the summary sheet states, e.g. 35.3, or None."""
    try:
        wb = _open(data)
        try:
            for row in wb[SUMMARY_SHEET].iter_rows(values_only=True):
                for v in row:
                    m = isinstance(v, str) and re.search(r"Diversification score:\s*([\d.]+)", v)
                    if m:
                        return float(m.group(1))
        finally:
            wb.close()
    except Exception:
        pass
    return None


def sections(data: bytes) -> list:
    """The analysis sheets in the workbook, in order (the summary and the explainers left out)."""
    try:
        wb = _open(data)
        try:
            return [n for n in wb.sheetnames
                    if n != SUMMARY_SHEET and not n.startswith(_META)]
        finally:
            wb.close()
    except Exception:
        return []
